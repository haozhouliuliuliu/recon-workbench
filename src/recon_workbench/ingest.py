"""输入读取：CSV / XLSX → 标准化表头后的原始行。

这一层只负责"把文件读成表"，不含任何业务判断；业务规则在 normalize.py。
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .errors import InputError
from .models import SourceName

HEADER_SCAN_LIMIT = 20
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}

REQUIRED_COLUMNS = ("invoice_no", "counterparty", "entry_date", "currency", "amount")

#: 中文/英文表头别名。比较前会做 NFKC、去空白、转小写、去下划线。
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_no": (
        "发票号",
        "发票号码",
        "单据号",
        "单据编号",
        "单号",
        "invoice_no",
        "invoice_no.",
        "invoice",
        "invoice number",
        "doc_no",
    ),
    "counterparty": (
        "供应商",
        "供应商名称",
        "往来单位",
        "客户",
        "客户名称",
        "公司",
        "counterparty",
        "supplier",
        "vendor",
    ),
    "entry_date": (
        "日期",
        "单据日期",
        "开票日期",
        "记账日期",
        "业务日期",
        "date",
        "entry_date",
        "invoice_date",
    ),
    "currency": ("币种", "货币", "币别", "currency", "ccy"),
    "amount": ("金额", "含税金额", "价税合计", "单据金额", "amount", "total", "amount_incl_tax"),
    "doc_type": ("单据类型", "类型", "票据类型", "doc_type", "type"),
    "note": ("备注", "摘要", "说明", "note", "remark", "memo"),
}

_ALIAS_TO_CANONICAL: dict[str, str] = {}
for _canonical, _aliases in COLUMN_ALIASES.items():
    for _alias in _aliases:
        _ALIAS_TO_CANONICAL[_alias] = _canonical


def _normalize_header(text: str) -> str:
    folded = unicodedata.normalize("NFKC", text).strip().casefold()
    return re.sub(r"[\s_\-]+", "", folded)


for _alias in list(_ALIAS_TO_CANONICAL):
    _ALIAS_TO_CANONICAL[_normalize_header(_alias)] = _ALIAS_TO_CANONICAL[_alias]


def canonical_column(header: str) -> str | None:
    """把任意中英文表头映射到规范列名；无法识别返回 None。"""

    return _ALIAS_TO_CANONICAL.get(_normalize_header(header))


@dataclass(frozen=True, slots=True)
class RawRow:
    row_number: int
    values: dict[str, str]


@dataclass(frozen=True, slots=True)
class RawTable:
    source: SourceName
    filename: str
    sha256: str
    header_row: int
    headers: tuple[str, ...]
    rows: tuple[RawRow, ...]


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _decode_csv(payload: bytes, path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise InputError(f"{path.name} 的编码无法识别（已尝试 UTF-8 / GBK），请先另存为 UTF-8 CSV 后重试")


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(line for line in text.splitlines()[:20] if line.strip())
    if not sample:
        return ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)).normalize(), "f")
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    return str(value).strip()


def _read_matrix_csv(payload: bytes, path: Path) -> list[list[str]]:
    text = _decode_csv(payload, path)
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    matrix: list[list[str]] = []
    for row in reader:
        matrix.append([cell.strip() for cell in row])
    return matrix


def _read_matrix_xlsx(path: Path, sheet: str | None) -> list[list[str]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 的异常类型较杂，统一收敛为输入错误
        raise InputError(f"{path.name} 不是可读取的 Excel 文件：{exc}") from exc

    try:
        if sheet is not None:
            if sheet not in workbook.sheetnames:
                raise InputError(f"{path.name} 中不存在工作表 {sheet!r}，可选：{', '.join(workbook.sheetnames)}")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]

        matrix: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            matrix.append([_cell_to_text(cell) for cell in row])
        return matrix
    finally:
        workbook.close()


def _locate_header(matrix: list[list[str]], path: Path) -> tuple[int, dict[str, int]]:
    """在前 N 行里找表头行，允许标题行、空行、说明行。

    返回 (表头行下标, 规范列名 → 列下标)。
    """

    for index, row in enumerate(matrix[:HEADER_SCAN_LIMIT]):
        mapping: dict[str, int] = {}
        duplicate: str | None = None
        for column_index, header in enumerate(row):
            canonical = canonical_column(header)
            if canonical is None:
                continue
            if canonical in mapping:
                duplicate = canonical
                break
            mapping[canonical] = column_index
        if duplicate is not None:
            raise InputError(f"{path.name} 第 {index + 1} 行表头重复：列 {duplicate!r} 出现了多次")
        missing = [name for name in REQUIRED_COLUMNS if name not in mapping]
        if not missing:
            return index, mapping

    preview = next((row for row in matrix[:5] if any(cell.strip() for cell in row)), [])
    found = ", ".join(header for header in preview if header.strip()) or "（空）"
    required_text = "、".join(REQUIRED_COLUMNS)
    raise InputError(
        f"{path.name} 找不到表头：前 {HEADER_SCAN_LIMIT} 行内必须同时包含 {required_text}。实际读到：{found}"
    )


def read_table(path: Path | str, source: SourceName, *, sheet: str | None = None) -> RawTable:
    """读取一份输入文件，返回带行号的原始表。"""

    input_path = Path(path)
    if not input_path.is_file():
        raise InputError(f"文件不存在：{input_path}")
    if input_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        supported = "、".join(sorted(SUPPORTED_SUFFIXES))
        raise InputError(f"不支持的文件类型 {input_path.suffix!r}，仅支持 {supported}")

    payload = input_path.read_bytes()
    if not payload:
        raise InputError(f"文件为空：{input_path.name}")

    if input_path.suffix.lower() == ".csv":
        matrix = _read_matrix_csv(payload, input_path)
    else:
        matrix = _read_matrix_xlsx(input_path, sheet)

    header_index, mapping = _locate_header(matrix, input_path)
    headers = tuple(canonical for _, canonical in sorted(mapping.items(), key=lambda kv: kv[1]))

    rows: list[RawRow] = []
    for index in range(header_index + 1, len(matrix)):
        raw_row = matrix[index]
        if not any(cell.strip() for cell in raw_row):
            continue
        values = {
            canonical: (raw_row[column] if column < len(raw_row) else "") for canonical, column in mapping.items()
        }
        rows.append(RawRow(row_number=index + 1, values=values))

    return RawTable(
        source=source,
        filename=input_path.name,
        sha256=sha256_bytes(payload),
        header_row=header_index + 1,
        headers=headers,
        rows=tuple(rows),
    )
