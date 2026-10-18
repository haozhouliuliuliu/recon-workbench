"""结果导出：规范化 JSON + 中文 Excel。

Excel 只做"给人看"的呈现，所有数值都能在 JSON 报告里找到同一份来源；
因此 Excel 的样式调整永远不会影响结论。
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Report, Status, money_str

STATUS_LABELS = {
    Status.MATCHED: "已匹配",
    Status.AMOUNT_MISMATCH: "金额不一致",
    Status.ONLY_IN_LEDGER: "仅台账有",
    Status.ONLY_IN_STATEMENT: "仅对账单有",
}

STATUS_FILLS = {
    Status.MATCHED: "E8F3E8",
    Status.AMOUNT_MISMATCH: "FBE4E4",
    Status.ONLY_IN_LEDGER: "FFF1DE",
    Status.ONLY_IN_STATEMENT: "FFF1DE",
}

MATCHED_BY_LABELS = {
    "invoice_no": "发票号",
    "date_amount": "日期+金额",
    None: "—",
}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="1F2937")
SUBTITLE_FONT = Font(name="微软雅黑", size=10, color="595959")
SECTION_FONT = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
BODY_FONT = Font(name="微软雅黑", size=10, color="262626")
MONEY_FORMAT = "#,##0.00"
ITEM_HEADERS = [
    "序号",
    "状态",
    "供应商",
    "发票号",
    "日期",
    "币种",
    "台账金额",
    "对账单金额",
    "差额",
    "匹配方式",
    "说明",
]


def _write_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=column, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 22


def _write_row(ws: Worksheet, row: int, values: list[object], *, money_columns: tuple[int, ...] = ()) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=column == len(values))
        if column in money_columns and isinstance(value, (int, float)):
            cell.number_format = MONEY_FORMAT
            cell.alignment = Alignment(horizontal="right", vertical="center")


def _auto_width(ws: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    for column, header in enumerate(headers, start=1):
        width = sum(2 if ord(char) > 127 else 1 for char in str(header)) + 4
        for row in rows[:200]:
            if column - 1 < len(row):
                text = "" if row[column - 1] is None else str(row[column - 1])
                width = max(width, min(60, sum(2 if ord(char) > 127 else 1 for char in text) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width


def _item_row(index: int, item) -> list[object]:
    ledger = item.ledger
    statement = item.statement
    primary = ledger or statement
    return [
        index,
        STATUS_LABELS[item.status],
        primary.counterparty,
        primary.invoice_no,
        primary.entry_date.isoformat(),
        primary.currency,
        float(ledger.amount) if ledger else None,
        float(statement.amount) if statement else None,
        float(item.delta) if item.delta is not None else None,
        MATCHED_BY_LABELS.get(item.matched_by, item.matched_by or "—"),
        item.reason,
    ]


def _write_items_sheet(ws: Worksheet, title: str, report: Report, items) -> None:
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ITEM_HEADERS))
    ws.cell(row=2, column=1, value=f"报告摘要 {report.digest}").font = SUBTITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(ITEM_HEADERS))

    header_row = 3
    _write_header(ws, header_row, ITEM_HEADERS)
    rows: list[list[object]] = []
    for index, item in enumerate(items, start=1):
        values = _item_row(index, item)
        rows.append(values)
        row = header_row + index
        _write_row(ws, row, values, money_columns=(7, 8, 9))
        fill = PatternFill("solid", fgColor=STATUS_FILLS[item.status])
        for column in (1, 2):
            ws.cell(row=row, column=column).fill = fill
    _auto_width(ws, ITEM_HEADERS, rows)
    ws.freeze_panes = f"A{header_row + 1}"
    if items:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(ITEM_HEADERS))}{header_row + len(items)}"


def _write_summary_sheet(ws: Worksheet, report: Report, generated_at: datetime) -> None:
    counts = report.counts
    ws.cell(row=1, column=1, value="企业对账核对报告").font = TITLE_FONT
    ws.merge_cells("A1:F1")
    ws.cell(row=2, column=1, value="本报告全部数据为虚构，仅用于流程演示").font = SUBTITLE_FONT
    ws.merge_cells("A2:F2")
    ws.cell(row=3, column=1, value=f"生成时间：{generated_at:%Y-%m-%d %H:%M:%S}").font = SUBTITLE_FONT
    ws.cell(row=4, column=1, value=f"报告摘要：{report.digest}").font = SUBTITLE_FONT

    ws.cell(row=6, column=1, value="一、结果概览").font = SECTION_FONT
    overview_headers = ["状态", "数量"]
    _write_header(ws, 7, overview_headers)
    overview_rows = [
        ["已匹配", counts.matched],
        ["金额不一致", counts.amount_mismatch],
        ["仅台账有", counts.only_in_ledger],
        ["仅对账单有", counts.only_in_statement],
        ["重复单据（已排除）", counts.duplicates],
        ["无效行（已排除）", counts.invalid_rows],
    ]
    for offset, values in enumerate(overview_rows, start=1):
        _write_row(ws, 7 + offset, values)

    start = 7 + len(overview_rows) + 3
    ws.cell(row=start, column=1, value="二、金额合计").font = SECTION_FONT
    totals_headers = ["币种", "台账合计", "对账单合计", "差额（台账-对账单）"]
    _write_header(ws, start + 1, totals_headers)
    for offset, totals in enumerate(report.totals, start=1):
        _write_row(
            ws,
            start + 1 + offset,
            [totals.currency, float(totals.ledger_total), float(totals.statement_total), float(totals.difference)],
            money_columns=(2, 3, 4),
        )

    start = start + len(report.totals) + 4
    ws.cell(row=start, column=1, value="三、输入文件").font = SECTION_FONT
    source_headers = ["来源", "文件名", "总行数", "有效行", "排除行", "SHA-256"]
    _write_header(ws, start + 1, source_headers)
    for offset, source in enumerate(report.sources, start=1):
        _write_row(
            ws,
            start + 1 + offset,
            [
                source.source.label_zh,
                source.filename,
                source.total_rows,
                source.accepted_rows,
                source.rejected_rows,
                source.sha256,
            ],
        )

    for column, width in enumerate([22, 26, 22, 22, 16, 66], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def _write_duplicates_sheet(ws: Worksheet, report: Report) -> None:
    headers = ["来源", "类型", "行号", "关联行号", "发票号", "供应商", "日期", "金额", "说明"]
    ws.cell(row=1, column=1, value="重复单据").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _write_header(ws, 3, headers)
    rows: list[list[object]] = []
    for offset, finding in enumerate(report.duplicates, start=1):
        entry = finding.entry
        values = [
            finding.source.label_zh,
            "完全重复" if finding.kind == "exact" else "疑似重复",
            entry.row_number,
            finding.related_row_number,
            entry.invoice_no,
            entry.counterparty,
            entry.entry_date.isoformat(),
            float(entry.amount),
            finding.reason,
        ]
        rows.append(values)
        _write_row(ws, 3 + offset, values, money_columns=(8,))
    _auto_width(ws, headers, rows)
    ws.freeze_panes = "A4"


def _write_invalid_sheet(ws: Worksheet, report: Report) -> None:
    headers = ["来源", "行号", "问题代码", "原始数据"]
    ws.cell(row=1, column=1, value="无效行").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _write_header(ws, 3, headers)
    rows: list[list[object]] = []
    for offset, row in enumerate(report.invalid_rows, start=1):
        values = [
            row.source.label_zh,
            row.row_number,
            "、".join(row.issue_codes),
            json.dumps(row.raw, ensure_ascii=False, sort_keys=True),
        ]
        rows.append(values)
        _write_row(ws, 3 + offset, values)
    _auto_width(ws, headers, rows)
    ws.freeze_panes = "A4"


def _write_issues_sheet(ws: Worksheet, report: Report) -> None:
    headers = ["来源", "行号", "级别", "代码", "字段", "说明"]
    ws.cell(row=1, column=1, value="问题清单").font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _write_header(ws, 3, headers)
    rows: list[list[object]] = []
    severity_label = {"error": "错误", "warning": "告警"}
    for offset, issue in enumerate(report.issues, start=1):
        values = [
            issue.source.label_zh,
            issue.row_number,
            severity_label[issue.severity.value],
            issue.code,
            issue.field or "",
            issue.message,
        ]
        rows.append(values)
        _write_row(ws, 3 + offset, values)
    _auto_width(ws, headers, rows)
    ws.freeze_panes = "A4"


def _write_meta_sheet(ws: Worksheet, report: Report, generated_at: datetime) -> None:
    headers = ["项目", "值"]
    ws.cell(row=1, column=1, value="运行信息").font = TITLE_FONT
    ws.merge_cells("A1:B1")
    _write_header(ws, 3, headers)
    policy = report.policy.model_dump(mode="json")
    rows = [
        ["工具版本", report.tool_version],
        ["报告 schema", report.schema_version],
        ["报告摘要（SHA-256）", report.digest],
        ["生成时间", generated_at.isoformat(sep=" ", timespec="seconds")],
        ["金额容差", money_str(report.policy.amount_tolerance)],
        ["允许币种", "、".join(report.policy.allowed_currencies)],
        ["策略快照", json.dumps(policy, ensure_ascii=False, sort_keys=True)],
    ]
    for offset, values in enumerate(rows, start=1):
        _write_row(ws, 3 + offset, values)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100


def write_excel(report: Report, path: Path | str, *, generated_at: datetime | None = None) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = generated_at or datetime.now()

    workbook = Workbook()
    workbook.properties.creator = "Recon Workbench"
    workbook.properties.title = "对账核对报告"
    workbook.properties.description = "全部数据为虚构，仅用于流程演示"

    summary = workbook.active
    summary.title = "汇总"
    _write_summary_sheet(summary, report, timestamp)

    _write_items_sheet(workbook.create_sheet("全部明细"), "全部核对明细", report, list(report.items))
    differences = [item for item in report.items if item.status is not Status.MATCHED]
    _write_items_sheet(workbook.create_sheet("差异明细"), f"差异明细（{len(differences)} 条）", report, differences)
    _write_duplicates_sheet(workbook.create_sheet("重复项"), report)
    _write_invalid_sheet(workbook.create_sheet("无效行"), report)
    _write_issues_sheet(workbook.create_sheet("问题清单"), report)
    _write_meta_sheet(workbook.create_sheet("运行信息"), report, timestamp)

    workbook.save(output_path)
    return output_path


def write_report_json(result_json: bytes, path: Path | str) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(result_json + b"\n")
    return output_path
