"""测试辅助：构造输入文件与领域对象。"""

from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from recon_workbench.ingest import RawRow, RawTable
from recon_workbench.models import Entry, SourceName

LEDGER_HEADERS = ["发票号", "供应商", "日期", "币种", "金额", "单据类型", "备注"]


def write_csv(path: Path, rows: list[list[str]], headers: list[str] | None = None, encoding: str = "utf-8-sig") -> Path:
    with path.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers or LEDGER_HEADERS)
        writer.writerows(rows)
    return path


def write_xlsx(
    path: Path,
    rows: list[list[object]],
    headers: list[str] | None = None,
    *,
    title: str | None = None,
    sheet: str = "对账单",
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    row_cursor = 1
    if title is not None:
        worksheet.cell(row=1, column=1, value=title)
        row_cursor = 3
    for column, header in enumerate(headers or LEDGER_HEADERS, start=1):
        worksheet.cell(row=row_cursor, column=column, value=header)
    for offset, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            worksheet.cell(row=row_cursor + offset, column=column, value=value)
    workbook.save(path)
    return path


def raw_table(
    rows: list[dict[str, str]],
    *,
    source: SourceName = SourceName.LEDGER,
    start_row: int = 2,
) -> RawTable:
    return RawTable(
        source=source,
        filename="input.csv",
        sha256="0" * 64,
        header_row=1,
        headers=tuple(sorted({key for row in rows for key in row})),
        rows=tuple(RawRow(row_number=start_row + index, values=row) for index, row in enumerate(rows)),
    )


def make_entry(
    *,
    source: SourceName = SourceName.LEDGER,
    row_number: int = 1,
    invoice: str = "INV-2099-0001",
    counterparty: str = "辰光金属（虚构）",
    day: date = date(2026, 8, 3),
    amount: str = "100.00",
    currency: str = "CNY",
    note: str = "",
) -> Entry:
    return Entry(
        entry_id=f"E{source.value[:1].upper()}{row_number:04d}",
        source=source,
        row_number=row_number,
        invoice_no=invoice,
        counterparty=counterparty,
        entry_date=day,
        currency=currency,
        amount=Decimal(amount),
        doc_type="送货单",
        note=note,
    )
