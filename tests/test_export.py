"""Excel 导出测试：工作表齐全、内容与报告一致、同参数重复导出取值一致。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from recon_workbench.export import write_excel
from recon_workbench.pipeline import run_reconciliation

SAMPLE_DIR = Path(__file__).resolve().parents[1] / "sample_data"
FIXED_TIME = datetime(2026, 10, 18, 12, 0, 0)


def _report():
    return run_reconciliation(SAMPLE_DIR / "ledger_2026_08.csv", SAMPLE_DIR / "statement_2026_08.xlsx").report


def _snapshot(path: Path) -> list[list[object]]:
    workbook = load_workbook(path)
    values: list[list[object]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            values.append(list(row))
    workbook.close()
    return values


def test_sheets_exist_in_order(tmp_path: Path) -> None:
    path = write_excel(_report(), tmp_path / "out.xlsx", generated_at=FIXED_TIME)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["汇总", "全部明细", "差异明细", "重复项", "无效行", "问题清单", "运行信息"]
    workbook.close()


def test_item_sheets_contain_expected_rows(tmp_path: Path) -> None:
    report = _report()
    path = write_excel(report, tmp_path / "out.xlsx", generated_at=FIXED_TIME)
    workbook = load_workbook(path)
    all_items = workbook["全部明细"]
    differences = workbook["差异明细"]
    assert all_items.max_row == 3 + len(report.items)
    assert (
        differences.max_row
        == 3 + report.counts.amount_mismatch + report.counts.only_in_ledger + report.counts.only_in_statement
    )
    workbook.close()


def test_summary_sheet_contains_totals(tmp_path: Path) -> None:
    path = write_excel(_report(), tmp_path / "out.xlsx", generated_at=FIXED_TIME)
    workbook = load_workbook(path)
    values = [cell for row in workbook["汇总"].iter_rows(values_only=True) for cell in row]
    assert "85480.75" not in values  # 金额以数值写入，不是字符串
    assert 85480.75 in values
    assert -6240.0 in values
    workbook.close()


def test_export_is_value_stable_across_runs(tmp_path: Path) -> None:
    first = write_excel(_report(), tmp_path / "a.xlsx", generated_at=FIXED_TIME)
    second = write_excel(_report(), tmp_path / "b.xlsx", generated_at=FIXED_TIME)
    assert _snapshot(first) == _snapshot(second)


def test_meta_sheet_has_digest(tmp_path: Path) -> None:
    report = _report()
    path = write_excel(report, tmp_path / "out.xlsx", generated_at=FIXED_TIME)
    workbook = load_workbook(path)
    values = [str(cell) for row in workbook["运行信息"].iter_rows(values_only=True) for cell in row if cell]
    assert report.digest in values
    assert "0.2.0" in values
    workbook.close()
